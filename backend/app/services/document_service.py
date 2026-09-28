import csv
import io
import json
from pathlib import Path
from typing import Optional
from app.core.logging_config import logger

class DocumentService:
    """Handles text extraction for documents: PDF, TXT, CSV, DOCX, JSON, Markdown."""

    def extract_text(self, file_path: Path, extension: str) -> str:
        ext = extension.lower().strip()
        try:
            if ext in (".txt", ".md"):
                return self._extract_plain_text(file_path)
            elif ext == ".csv":
                return self._extract_csv(file_path)
            elif ext == ".json":
                return self._extract_json(file_path)
            elif ext == ".pdf":
                return self._extract_pdf(file_path)
            elif ext == ".docx":
                return self._extract_docx(file_path)
            else:
                return ""
        except Exception as e:
            logger.error(f"Failed to extract document text from {file_path.name}: {e}")
            return f"[Notice: Content from {file_path.name} could not be extracted: {str(e)}]"

    def _extract_plain_text(self, file_path: Path) -> str:
        data = file_path.read_bytes()
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")

    def _extract_csv(self, file_path: Path, max_rows: int = 200) -> str:
        text = self._extract_plain_text(file_path)
        reader = csv.reader(io.StringIO(text))
        rows = []
        for i, row in enumerate(reader):
            if i >= max_rows:
                rows.append(f"... [Truncated after {max_rows} rows]")
                break
            rows.append(" | ".join([cell.strip() for cell in row]))
        return "\n".join(rows)

    def _extract_json(self, file_path: Path) -> str:
        text = self._extract_plain_text(file_path)
        try:
            parsed = json.loads(text)
            return json.dumps(parsed, indent=2, ensure_ascii=False)
        except Exception:
            return text

    def _extract_pdf(self, file_path: Path, max_pages: int = 50) -> str:
        text_parts = []
        try:
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            total_pages = len(reader.pages)
            pages_to_read = min(total_pages, max_pages)
            for page_num in range(pages_to_read):
                page = reader.pages[page_num]
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(f"--- Page {page_num + 1} ---\n{page_text.strip()}")
            if total_pages > max_pages:
                text_parts.append(f"... [Truncated after {max_pages} pages of {total_pages}]")
        except Exception as e:
            # Fallback to PyMuPDF if available
            try:
                import fitz
                doc = fitz.open(str(file_path))
                total_pages = len(doc)
                pages_to_read = min(total_pages, max_pages)
                for page_num in range(pages_to_read):
                    page_text = doc[page_num].get_text() or ""
                    if page_text.strip():
                        text_parts.append(f"--- Page {page_num + 1} ---\n{page_text.strip()}")
                if total_pages > max_pages:
                    text_parts.append(f"... [Truncated after {max_pages} pages of {total_pages}]")
            except Exception as inner_e:
                raise RuntimeError(f"PDF extraction error: {e} | {inner_e}")

        extracted = "\n\n".join(text_parts).strip()
        return extracted or "[Notice: PDF contains no readable text, scanned images, or empty pages.]"

    def _extract_docx(self, file_path: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(file_path))
            parts = []
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    parts.append(t)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    if any(row_cells):
                        parts.append(" | ".join(row_cells))
            return "\n".join(parts)
        except Exception as e:
            raise RuntimeError(f"DOCX extraction error: {e}")

document_service = DocumentService()
