import io
import pytest
from pathlib import Path
from fastapi import UploadFile, HTTPException
from app.utils.file_utils import validate_file_metadata, MAX_FILE_SIZE_BYTES
from app.services.document_service import document_service
from app.services.image_service import image_service
from app.services.file_service import FileService
from app.services.chat_service import ChatService
from app.services.llm_service import MockLLMProvider, LLMService

def test_file_metadata_validation():
    # Valid types
    name, ext = validate_file_metadata("report.pdf", 1000)
    assert ext == ".pdf"
    assert name == "report.pdf"

    name, ext = validate_file_metadata("photo.png", 500)
    assert ext == ".png"

    # Unsupported type
    with pytest.raises(ValueError) as exc:
        validate_file_metadata("script.exe", 1000)
    assert "Unsupported file type" in str(exc.value)

    # Oversized file
    with pytest.raises(ValueError) as exc:
        validate_file_metadata("big.pdf", MAX_FILE_SIZE_BYTES + 1)
    assert "exceeds the 10 MB limit" in str(exc.value)

def test_document_extraction_txt_csv_json(tmp_path):
    # TXT
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Hello world, this is a plain text file.", encoding="utf-8")
    assert "Hello world" in document_service.extract_text(txt_file, ".txt")

    # CSV
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("name,role,level\nSanju,AI Engineer,Senior\nAlice,Developer,Mid", encoding="utf-8")
    csv_text = document_service.extract_text(csv_file, ".csv")
    assert "Sanju | AI Engineer | Senior" in csv_text

    # JSON
    json_file = tmp_path / "config.json"
    json_file.write_text('{"project": "MemoryBot", "version": "2.0"}', encoding="utf-8")
    json_text = document_service.extract_text(json_file, ".json")
    assert '"project": "MemoryBot"' in json_text

def test_document_extraction_pdf(tmp_path):
    import pypdf
    pdf_file = tmp_path / "sample.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with open(pdf_file, "wb") as f:
        writer.write(f)

    extracted = document_service.extract_text(pdf_file, ".pdf")
    assert isinstance(extracted, str)

def test_document_extraction_docx(tmp_path):
    import docx
    doc = docx.Document()
    doc.add_heading("MemoryBot Documentation", level=1)
    doc.add_paragraph("This is a test paragraph inside a docx document.")
    docx_file = tmp_path / "test.docx"
    doc.save(str(docx_file))

    extracted = document_service.extract_text(docx_file, ".docx")
    assert "MemoryBot Documentation" in extracted
    assert "test paragraph inside a docx document" in extracted

@pytest.mark.asyncio
async def test_file_upload_and_chat_flow(test_db, test_user, tmp_path):
    file_service = FileService(test_db)
    mock_llm = LLMService(provider=MockLLMProvider())
    chat_service = ChatService(test_db, llm_service=mock_llm)

    # 1. Upload text document
    content = b"Python Architecture and Concepts:\nConcept 1: Variables\nConcept 2: Loops"
    upload = UploadFile(filename="python_guide.txt", file=io.BytesIO(content))
    upload_resp = await file_service.save_uploaded_file(upload, user_id=test_user.id)
    assert upload_resp.id is not None
    assert upload_resp.filename == "python_guide.txt"

    # 2. First chat turn with attachment
    chat_resp1 = chat_service.process_chat_message(
        user_id=test_user.id,
        message="What is this document about?",
        attachment_ids=[upload_resp.id]
    )
    assert chat_resp1.conversation_id is not None
    assert len(chat_resp1.user_message.attachments) == 1
    assert chat_resp1.user_message.attachments[0].filename == "python_guide.txt"
    assert "document provides an overview of Python" in chat_resp1.assistant_message.content

    # 3. Second chat turn in SAME conversation asking follow-up without re-uploading file
    chat_resp2 = chat_service.process_chat_message(
        user_id=test_user.id,
        message="What is the first concept?",
        conversation_id=chat_resp1.conversation_id
    )
    assert "first concept is Python Basics" in chat_resp2.assistant_message.content

    # 4. Third chat turn asking for example
    chat_resp3 = chat_service.process_chat_message(
        user_id=test_user.id,
        message="Explain that with an example.",
        conversation_id=chat_resp1.conversation_id
    )
    assert "Here is an example" in chat_resp3.assistant_message.content

@pytest.mark.asyncio
async def test_multimodal_image_upload_and_chat(test_db, test_user):
    from PIL import Image
    file_service = FileService(test_db)
    mock_llm = LLMService(provider=MockLLMProvider())
    chat_service = ChatService(test_db, llm_service=mock_llm)

    # Create dummy PNG image
    img = Image.new("RGB", (100, 100), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    upload = UploadFile(filename="architecture.png", file=buf)
    upload_resp = await file_service.save_uploaded_file(upload, user_id=test_user.id)
    assert upload_resp.filename == "architecture.png"

    chat_resp = chat_service.process_chat_message(
        user_id=test_user.id,
        message="What is shown in this image?",
        attachment_ids=[upload_resp.id]
    )
    assert "architectural diagram" in chat_resp.assistant_message.content
